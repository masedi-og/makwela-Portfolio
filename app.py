import csv
import json
import os
import re
import shutil
import threading
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook


REQUIRED_COLUMNS = [
    "Supplier",
    "Supplier Code",
    "Catalogue Code",
    "Product Description",
    "Parent Code",
    "Range 1 Name",
    "Range 1 Number",
    "Missing Product Image",
    "Missing Parent Image",
    "Missing Range 1 Image",
    "Comments",
]


@dataclass
class ResultRow:
    row_number: int
    supplier: str
    supplier_code: str
    catalogue_code: str
    process_type: str
    status: str
    found_in: str
    match_type: str
    found_files: str
    output_files: str
    notes: str


@dataclass
class WorkItem:
    row_number: int
    supplier: str
    supplier_code: str
    catalogue_code: str
    process_type: str


class ImageFinder:
    def __init__(self, supplier_config: Dict):
        self.supplier_config = supplier_config

    @staticmethod
    def _is_tiff(path: Path) -> bool:
        return path.suffix.lower() in {".tif", ".tiff"}

    @staticmethod
    def _token_contains(stem: str, code: str) -> bool:
        pattern = re.compile(rf"(^|[ _-]){re.escape(code)}($|[ _-])", re.IGNORECASE)
        return bool(pattern.search(stem))

    def _collect_matches_in_folder(self, folder: Path, code: str) -> List[Path]:
        if not folder.exists() or not folder.is_dir():
            return []
        matches = []
        for child in folder.iterdir():
            if not child.is_file() or not self._is_tiff(child):
                continue
            stem = child.stem
            if stem.lower().startswith(code.lower()) or self._token_contains(stem, code):
                matches.append(child)
        return matches

    def find_candidate_set(self, code: str) -> Tuple[str, str, List[Path], str]:
        pre = Path(self.supplier_config["source"]["pre_edit"])
        master = Path(self.supplier_config["source"]["master"])

        candidates = []
        for source_name, folder in (("PreEdit", pre), ("Master", master)):
            matches = self._collect_matches_in_folder(folder, code)
            if not matches:
                continue

            exact_primary = None
            secondaries = []
            for file in matches:
                stem = file.stem
                if stem.lower() == code.lower():
                    exact_primary = file
                sec_match = re.match(rf"^{re.escape(code)}_(\d+)$", stem, flags=re.IGNORECASE)
                if sec_match:
                    idx = int(sec_match.group(1))
                    if idx >= 2:
                        secondaries.append((idx, file))

            if exact_primary:
                secondaries.sort(key=lambda x: x[0])
                files = [exact_primary] + [f for _, f in secondaries]
                candidates.append((source_name, "StartsWith", files, "Exact primary found"))
                continue

            primaries = []
            for file in matches:
                stem = file.stem
                if re.match(rf"^{re.escape(code)}_1$", stem, flags=re.IGNORECASE):
                    continue
                if re.match(rf"^{re.escape(code)}_(\d+)$", stem, flags=re.IGNORECASE):
                    continue
                if stem.lower().startswith(code.lower()):
                    mtype = "StartsWith"
                else:
                    mtype = "TokenContains"
                primaries.append((file, mtype))

            if len(primaries) == 1:
                candidates.append((source_name, primaries[0][1], [primaries[0][0]], "Single non-exact candidate"))
            elif len(primaries) > 1:
                candidates.append((source_name, "Mixed", [p[0] for p in primaries], "Multiple primaries in folder"))

        if not candidates:
            return "Missing", "", [], "No matching TIFF files found"

        if len(candidates) > 1:
            all_files = []
            for _, _, files, _ in candidates:
                all_files.extend(files)
            return "Ambiguous", "", all_files, "Multiple candidate sets found across sources"

        found_in, match_type, files, note = candidates[0]
        if note.startswith("Multiple"):
            return "Ambiguous", found_in, files, note
        return "Success", found_in, files, note


class DesktopAgentApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Desktop Agent MVP")
        self.root.geometry("1200x760")

        self.config = self._load_config()
        self.excel_path = tk.StringVar()
        self.dry_run = tk.BooleanVar(value=True)
        self.allow_overwrite = tk.BooleanVar(value=False)

        self.results: List[ResultRow] = []

        self._build_ui()

    def _load_config(self) -> Dict:
        config_file = Path("config.json")
        if not config_file.exists():
            raise FileNotFoundError("config.json not found")
        with open(config_file, "r", encoding="utf-8") as f:
            return json.load(f)

    def _build_ui(self):
        top = ttk.Frame(self.root, padding=10)
        top.pack(fill="x")

        ttk.Label(top, text="Excel File:").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.excel_path, width=90).grid(row=0, column=1, sticky="we", padx=5)
        ttk.Button(top, text="Browse", command=self.browse_excel).grid(row=0, column=2)
        top.columnconfigure(1, weight=1)

        supplier_frame = ttk.LabelFrame(self.root, text="Suppliers", padding=10)
        supplier_frame.pack(fill="x", padx=10, pady=6)

        self.supplier_listbox = tk.Listbox(supplier_frame, selectmode=tk.MULTIPLE, height=6, exportselection=False)
        self.supplier_listbox.pack(fill="x")
        for supplier in self.config.get("suppliers", {}).keys():
            self.supplier_listbox.insert(tk.END, supplier)

        options = ttk.Frame(self.root, padding=10)
        options.pack(fill="x")

        ttk.Checkbutton(options, text="Dry run", variable=self.dry_run).pack(side="left", padx=5)
        ttk.Checkbutton(options, text="Allow overwrite", variable=self.allow_overwrite).pack(side="left", padx=5)
        ttk.Button(options, text="Run", command=self.run).pack(side="right", padx=5)

        self.progress = ttk.Progressbar(self.root, orient="horizontal", mode="determinate")
        self.progress.pack(fill="x", padx=10, pady=5)

        log_frame = ttk.LabelFrame(self.root, text="Live Log", padding=8)
        log_frame.pack(fill="both", expand=False, padx=10, pady=5)
        self.log = tk.Text(log_frame, height=9)
        self.log.pack(fill="both", expand=True)

        result_frame = ttk.LabelFrame(self.root, text="Results", padding=8)
        result_frame.pack(fill="both", expand=True, padx=10, pady=5)

        columns = [
            "row_number", "supplier", "catalogue_code", "process_type", "status",
            "found_in", "match_type", "found_files", "output_files", "notes"
        ]
        self.tree = ttk.Treeview(result_frame, columns=columns, show="headings")
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=120)
        self.tree.pack(fill="both", expand=True)

        export_frame = ttk.Frame(self.root, padding=10)
        export_frame.pack(fill="x")
        ttk.Button(export_frame, text="Export CSV", command=self.export_csv).pack(side="left", padx=5)
        ttk.Button(export_frame, text="Export XLSX", command=self.export_xlsx).pack(side="left", padx=5)

    def browse_excel(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.excel_path.set(path)

    def selected_suppliers(self) -> List[str]:
        indices = self.supplier_listbox.curselection()
        return [self.supplier_listbox.get(i) for i in indices]

    def append_log(self, message: str):
        self.log.insert(tk.END, message + "\n")
        self.log.see(tk.END)
        self.root.update_idletasks()

    def run(self):
        if not self.excel_path.get():
            messagebox.showerror("Error", "Please select an Excel file")
            return
        suppliers = self.selected_suppliers()
        if not suppliers:
            messagebox.showerror("Error", "Please select at least one supplier")
            return

        thread = threading.Thread(target=self._run_worker, args=(suppliers,), daemon=True)
        thread.start()

    def _run_worker(self, suppliers: List[str]):
        try:
            rows = self._load_rows(self.excel_path.get())
        except Exception as exc:
            messagebox.showerror("Error", f"Failed loading Excel: {exc}")
            return

        work_items = self._build_work_items(rows, suppliers)
        self.results.clear()
        for item in self.tree.get_children():
            self.tree.delete(item)

        total = max(1, len(work_items))
        self.progress["value"] = 0
        self.progress["maximum"] = total

        for index, work in enumerate(work_items, start=1):
            result = self._process_item(work)
            self.results.append(result)
            self.tree.insert("", tk.END, values=(
                result.row_number,
                result.supplier,
                result.catalogue_code,
                result.process_type,
                result.status,
                result.found_in,
                result.match_type,
                result.found_files,
                result.output_files,
                result.notes,
            ))
            self.progress["value"] = index
            self.append_log(
                f"Row {result.row_number} | {result.process_type} | {result.catalogue_code} => {result.status}"
            )

        self._log_missing_grouped_by_supplier()
        self.append_log("Run completed.")

    def _load_rows(self, xlsx_path: str) -> List[Dict]:
        wb = load_workbook(xlsx_path)
        ws = wb.active
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

        missing_cols = [col for col in REQUIRED_COLUMNS if col not in headers]
        if missing_cols:
            raise ValueError(f"Missing required columns: {missing_cols}")

        idx_map = {name: headers.index(name) for name in REQUIRED_COLUMNS}
        data = []
        for r in range(2, ws.max_row + 1):
            row = {}
            empty = True
            for col in REQUIRED_COLUMNS:
                val = ws.cell(row=r, column=idx_map[col] + 1).value
                row[col] = "" if val is None else str(val).strip()
                if row[col]:
                    empty = False
            if not empty:
                row["__row_number"] = r
                data.append(row)
        return data

    def _build_work_items(self, rows: List[Dict], suppliers: List[str]) -> List[WorkItem]:
        items: List[WorkItem] = []
        for row in rows:
            if row["Supplier"] not in suppliers:
                continue
            code = row["Catalogue Code"]
            if not code:
                continue

            if row["Missing Product Image"].lower() == "yes":
                items.append(WorkItem(row["__row_number"], row["Supplier"], row["Supplier Code"], code, "CHILD"))
            if row["Missing Parent Image"].lower() == "yes":
                items.append(WorkItem(row["__row_number"], row["Supplier"], row["Supplier Code"], code, "PARENT"))
            if row["Missing Range 1 Image"].lower() == "yes":
                if row["Range 1 Name"] or row["Range 1 Number"]:
                    items.append(WorkItem(row["__row_number"], row["Supplier"], row["Supplier Code"], code, "RANGE"))
        return items

    def _destination_tiff_working(self, supplier: str, process_type: str) -> Path:
        cfg = self.config["suppliers"][supplier]
        key_map = {"CHILD": "child", "PARENT": "parent", "RANGE": "range"}
        base = Path(cfg["upload"][key_map[process_type]])
        return base / "TIFF_WORKING"

    def _process_item(self, item: WorkItem) -> ResultRow:
        supplier_cfg = self.config["suppliers"].get(item.supplier)
        if not supplier_cfg:
            return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                             "Missing", "", "", "", "", "Supplier missing in config")

        finder = ImageFinder(supplier_cfg)
        status, found_in, match_type, note = "Missing", "", "", ""
        files: List[Path] = []

        status, found_in, files, note = finder.find_candidate_set(item.catalogue_code)
        if status == "Missing":
            return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                             status, found_in, match_type, "", "", note)
        if status == "Ambiguous":
            return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                             status, found_in, "Mixed", "; ".join(str(f) for f in files), "", note)

        exact = any(f.stem.lower() == item.catalogue_code.lower() for f in files)
        match_type = "StartsWith" if exact else "TokenContains"

        destination = self._destination_tiff_working(item.supplier, item.process_type)
        output_files: List[str] = []

        try:
            if not self.dry_run.get():
                destination.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                             "Missing", found_in, match_type, "; ".join(str(f) for f in files), "", f"Create destination failed: {exc}")

        for idx, src in enumerate(files):
            out_name = f"{item.catalogue_code}.tif" if idx == 0 else f"{item.catalogue_code}_{idx+1}.tif"
            dst = destination / out_name
            output_files.append(str(dst))

            if self.dry_run.get():
                continue

            if dst.exists() and not self.allow_overwrite.get():
                return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                                 "Missing", found_in, match_type, "; ".join(str(f) for f in files),
                                 "; ".join(output_files), f"Destination exists and overwrite disabled: {dst}")
            shutil.copy2(src, dst)

        return ResultRow(item.row_number, item.supplier, item.supplier_code, item.catalogue_code, item.process_type,
                         "Success", found_in, match_type, "; ".join(str(f) for f in files), "; ".join(output_files), note)

    def _log_missing_grouped_by_supplier(self):
        grouped: Dict[str, List[ResultRow]] = {}
        for r in self.results:
            if r.status == "Missing":
                grouped.setdefault(r.supplier, []).append(r)

        if not grouped:
            self.append_log("No missing request list entries.")
            return

        self.append_log("Missing request list grouped by supplier:")
        for supplier, rows in grouped.items():
            codes = ", ".join(sorted({row.catalogue_code for row in rows}))
            self.append_log(f"- {supplier}: {codes}")

    def export_csv(self):
        if not self.results:
            messagebox.showinfo("Info", "No results to export")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV", "*.csv")])
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(asdict(self.results[0]).keys()))
            writer.writeheader()
            for row in self.results:
                writer.writerow(asdict(row))
        messagebox.showinfo("Done", "CSV exported")

    def export_xlsx(self):
        if not self.results:
            messagebox.showinfo("Info", "No results to export")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        headers = list(asdict(self.results[0]).keys())
        ws.append(headers)
        for row in self.results:
            ws.append([asdict(row)[h] for h in headers])
        wb.save(path)
        messagebox.showinfo("Done", "XLSX exported")


if __name__ == "__main__":
    root = tk.Tk()
    app = DesktopAgentApp(root)
    root.mainloop()
